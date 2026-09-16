# -*- coding: utf-8 -*-
"""蓝色科技风 BIS 时间同步 Demo 演示 PPT（带图形化示意图）。"""

from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

OUT = "/home/devin/ncs/lim_work/iso_time_sync/BIS时间同步Demo演示.pptx"

# ---- 配色（蓝色科技风）----
NAVY   = (0x0A, 0x25, 0x40)
NAVY2  = (0x0E, 0x30, 0x52)
BLUE   = (0x1F, 0x6F, 0xEB)
BLUE2  = (0x2E, 0x86, 0xFF)
CYAN   = (0x00, 0xD4, 0xFF)
LIGHT  = (0xEA, 0xF3, 0xFF)
WHITE  = (0xFF, 0xFF, 0xFF)
INK    = (0x1A, 0x2B, 0x3C)
GRAY   = (0x5A, 0x6B, 0x7B)
CARD   = (0xF5, 0xF9, 0xFF)
IMG    = "/home/devin/ncs/lim_work/iso_time_sync/scene_images"

prs = Presentation()
prs.slide_width = I(13.333)
prs.slide_height = I(7.5)
BLANK = prs.slide_layouts[6]

def set_font(run, size=18, color=INK, bold=False, name="微软雅黑"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = C(*color)
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', name)

def add_rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = C(*fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = C(*line)
        sp.line.width = Pt(1.2)
    sp.shadow.inherit = False
    return sp

def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name="微软雅黑"):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        set_font(r, size=size, color=color, bold=bold, name=name)
    return tb

def add_shape_text(sp, text, size=15, color=WHITE, bold=True, name="微软雅黑"):
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(6); tf.margin_right = Pt(6)
    tf.margin_top = Pt(3); tf.margin_bottom = Pt(3)
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = ln
        set_font(r, size=size, color=color, bold=bold, name=name)

def add_arrow(slide, x1, y1, x2, y2, color=BLUE, width=2.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = C(*color)
    c.line.width = Pt(width)
    c.shadow.inherit = False
    return c

def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = C(*color)

def deco_dots(slide, n=6):
    import random
    for _ in range(n):
        x = random.uniform(0.3, 12.5); y = random.uniform(0.4, 7.0)
        d = random.uniform(0.06, 0.22)
        add_rect(slide, x, y, d, d, CYAN, shape=MSO_SHAPE.OVAL)

def title_bar(slide, title, subtitle=None):
    add_rect(slide, 0, 0, 13.333, 1.15, NAVY)
    add_rect(slide, 0.55, 1.0, 2.2, 0.12, CYAN)
    add_text(slide, 0.55, 0.18, 12.2, 0.8, title, 28, WHITE, True)
    if subtitle:
        add_text(slide, 0.55, 1.28, 12.2, 0.5, subtitle, 14, GRAY)

# ================= 1 封面 =================
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
deco_dots(s, 6)
add_text(s, 0.7, 1.7, 6.0, 1.5, "nRF54L15\n蓝牙 BIS 时间同步", 40, WHITE, True)
add_text(s, 0.7, 3.5, 6.0, 1.0, "机器人大脑与四肢关节\n时钟同步 · 技术方案与演示", 22, CYAN)
add_rect(s, 0.75, 4.85, 2.6, 0.06, CYAN, shape=MSO_SHAPE.RECTANGLE)
tags = ["BLE 5.2 ISO", "1 发 2 收", "1µs 精度", "200Hz 更新"]
x = 0.7
for t in tags:
    add_shape_text(add_rect(s, x, 5.15, 1.75, 0.55, NAVY2, line=CYAN), t, 13, CYAN)
    x += 1.95
add_text(s, 0.7, 6.7, 6.0, 0.5, "Nordic nRF Connect SDK v3.4.0 · nRF54L15", 12, GRAY)
add_rect(s, 6.9, 0.9, 6.2, 4.35, None, line=CYAN)
s.shapes.add_picture(f"{IMG}/Brain_to_joint_clock_synchroni_2026-09-15T04-52-51.png", I(7.0), I(1.0), width=I(6.0))

# ================= 2 目录 =================
s = prs.slides.add_slide(BLANK)
bg(s, LIGHT)
title_bar(s, "目录", "CONTENTS")
items = [
    ("01", "演示目标", "微秒级时间同步的目标与关键指标"),
    ("02", "技术原理", "时间戳载体 + 时钟伺服 + 定时呈现（一页看懂）"),
    ("03", "系统架构与演示环境", "1 发 2 收架构 · 软硬件与工具"),
    ("04", "实测数据与优势", "0µs 误差 · 应用场景"),
]
y = 1.7
for num, t, d in items:
    sp = add_rect(s, 0.8, y, 0.9, 0.9, BLUE, shape=MSO_SHAPE.OVAL)
    add_shape_text(sp, num, 20, WHITE)
    add_text(s, 2.0, y, 10.5, 0.5, t, 20, INK, True)
    add_text(s, 2.0, y + 0.45, 10.5, 0.4, d, 14, GRAY)
    y += 1.35

# ================= 3 演示目标 =================
s = prs.slides.add_slide(BLANK)
bg(s, LIGHT)
title_bar(s, "演示目标", "OBJECTIVE")
goals = [
    ("多节点统一时间基准", "1 个发送端 + 2 个接收端，长期共享同一时间基准"),
    ("微秒级打戳一致", "三方对同一事件打出相同时间戳，精度约 1µs"),
    ("硬件级确定性呈现", "P1.10 输出 1µs 窄脉冲，边沿对齐可被逻辑分析仪直观观测"),
]
yy = 1.7
for t, d in goals:
    add_rect(s, 0.7, yy, 6.6, 1.55, WHITE, line=BLUE2)
    add_rect(s, 0.7, yy, 0.16, 1.55, CYAN, shape=MSO_SHAPE.RECTANGLE)
    add_text(s, 1.05, yy + 0.18, 6.0, 0.5, t, 18, NAVY, True)
    add_text(s, 1.05, yy + 0.68, 6.0, 0.75, d, 13, GRAY)
    yy += 1.75
add_rect(s, 7.6, 1.6, 5.3, 5.3, None, line=BLUE2)
s.shapes.add_picture(f"{IMG}/Robot_brain_controller_alignin_2026-09-15T04-52-51.png", I(7.7), I(1.7), width=I(5.1))

# ================= 4 什么是 BIS =================
s = prs.slides.add_slide(BLANK)
bg(s, LIGHT)
title_bar(s, "什么是 BLE BIS", "BROADCAST ISOCHRONOUS STREAM · 规范解读")

# 左：定义与规范
add_rect(s, 0.7, 1.6, 5.9, 5.35, WHITE, line=BLUE)
add_rect(s, 0.7, 1.6, 5.9, 0.6, BLUE, shape=MSO_SHAPE.RECTANGLE)
add_shape_text(add_rect(s, 0.7, 1.6, 5.9, 0.6, BLUE, shape=MSO_SHAPE.RECTANGLE), "定义与规范依据", 14, WHITE)
add_text(s, 0.95, 2.35, 5.4, 4.4,
    "• BIS = Broadcast Isochronous Stream\n"
    "  广播等时流，BLE 5.2 引入的\n"
    "  LE Isochronous Channels（LE-ISO）特性\n\n"
    "• 多条 BIS 组成一个 BIG\n"
    "  （Broadcast Isochronous Group），\n"
    "  一对多广播，无需连接/回传\n\n"
    "• 规范：Bluetooth Core Spec\n"
    "  Vol 6 Part G（Isochronous\n"
    "  Adaptation Layer，ISOAL）\n\n"
    "• 核心机制：以固定 ISO_Interval 发送 SDU，\n"
    "  BIG 锚点作为同步参考时间基准", 12.5, INK)

# 右：关键参数 + 同步流程
add_rect(s, 6.85, 1.6, 5.85, 5.35, WHITE, line=BLUE2)
add_rect(s, 6.85, 1.6, 5.85, 0.6, BLUE2, shape=MSO_SHAPE.RECTANGLE)
add_shape_text(add_rect(s, 6.85, 1.6, 5.85, 0.6, BLUE2, shape=MSO_SHAPE.RECTANGLE), "关键参数与同步流程", 14, WHITE)
add_text(s, 7.1, 2.3, 5.4, 1.3,
    "关键参数：ISO_Interval（5ms~4s，\n"
    "步进 1.25ms）· SDU_Interval · RTN · NSE\n"
    "Max Transport Latency · 2M PHY", 11.5, INK)
steps = [
    "Broadcaster 发 periodic adv\n（携带 BIGInfo 控制信息）",
    "Sync Receiver 扫描并\n同步 Periodic Advertising",
    "按 BIGInfo 建立 BIG 同步",
    "定时接收 SDU，BIG 锚点\n作为时间基准",
]
y = 3.8
for i, t in enumerate(steps):
    add_shape_text(add_rect(s, 7.1, y, 0.5, 0.5, BLUE2, shape=MSO_SHAPE.OVAL), str(i+1), 11, WHITE)
    add_text(s, 7.75, y - 0.05, 4.8, 0.6, t, 10.5, INK)
    y += 0.63

# ================= 5 技术原理（一页看懂） =================
s = prs.slides.add_slide(BLANK)
bg(s, LIGHT)
title_bar(s, "技术原理", "PRINCIPLE · 一页看懂")
# 主流程
tx = add_rect(s, 0.8, 1.55, 3.6, 1.15, BLUE)
add_shape_text(tx, "发送端 TX\nSDU = [trigger][counter][tx_ts]", 13, WHITE)
rx = add_rect(s, 9.0, 1.55, 3.6, 1.15, BLUE2)
add_shape_text(rx, "接收端 RX\n解析 → 时钟伺服 → shared_ts", 13, WHITE)
add_arrow(s, 4.4, 2.12, 9.0, 2.12, CYAN, 3)
add_text(s, 4.3, 1.3, 5.0, 0.5, "BIS 广播（5ms 间隔）", 12, GRAY, align=PP_ALIGN.CENTER)
# 三个原理卡片
cards = [
    ("① 时间戳载体", "发送端把 SDU 同步参考时间戳\ntx_ts 写入 payload；接收端拿到\n同一瞬间的本地读数 info->ts",
     "参考对 (tx_ts, info->ts)"),
    ("② 时钟伺服", "offset = 发送端时钟 − 接收端时钟\n中位数估计 + ppm 漂移估计\n外推到当前时刻",
     "shared = local + offset + drift×age"),
    ("③ 定时呈现", "GRTC → DPPI → GPIOTE 硬件链路\n在 presentation 时刻驱动 P1.10\n免软件线程抖动",
     "GRTC → DPPI → GPIOTE"),
]
x = 0.8
for t, d, f in cards:
    add_rect(s, x, 3.15, 3.8, 3.15, WHITE, line=BLUE)
    add_rect(s, x, 3.15, 3.8, 0.55, BLUE, shape=MSO_SHAPE.RECTANGLE)
    add_shape_text(add_rect(s, x, 3.15, 3.8, 0.55, BLUE, shape=MSO_SHAPE.RECTANGLE), t, 14, WHITE)
    add_text(s, x + 0.25, 3.85, 3.3, 1.55, d, 11.5, INK)
    add_text(s, x + 0.25, 5.55, 3.3, 0.6, f, 12, BLUE2, True)
    x += 4.1
add_text(s, 0.8, 6.55, 11.7, 0.5, "关键：tx_ts 与 info->ts 是同一物理瞬间（BIG 锚点）在两个自由运行时钟下的读数 → 据此估计偏移与漂移", 12, GRAY, align=PP_ALIGN.CENTER)

# ================= 6 系统架构 =================
s = prs.slides.add_slide(BLANK)
bg(s, LIGHT)
title_bar(s, "系统架构", "ARCHITECTURE · 1 发 2 收")
add_shape_text(add_rect(s, 0.7, 1.7, 2.6, 0.7, BLUE), "发送端 TX", 15, WHITE)
y = 2.6
for t in ["触发事件生成", "组包(嵌入时间戳)", "BIS 广播", "本地定时呈现"]:
    add_shape_text(add_rect(s, 0.7, y, 2.6, 0.75, WHITE, line=BLUE2), t, 13, NAVY)
    y += 0.95
add_arrow(s, 3.4, 4.3, 5.3, 4.3, CYAN, 3)
add_text(s, 3.3, 4.5, 2.2, 0.5, "BIS 广播\n(5ms 间隔)", 12, GRAY, align=PP_ALIGN.CENTER)
for cx in [5.6, 9.3]:
    add_shape_text(add_rect(s, cx, 1.7, 3.2, 0.7, BLUE2), "接收端 RX", 15, WHITE)
    yy = 2.6
    for t in ["解析 SDU", "时钟伺服(offset/drift)", "shared_ts 打戳", "P1.10 定时呈现"]:
        add_shape_text(add_rect(s, cx, yy, 3.2, 0.75, WHITE, line=BLUE2), t, 13, NAVY)
        yy += 0.95

# ================= 7 Demo 环境 =================
s = prs.slides.add_slide(BLANK)
bg(s, LIGHT)
title_bar(s, "Demo 环境", "SOFTWARE · HARDWARE · TOOLS")
cols = [
    ("硬件 Hardware", BLUE, [
        "nRF54L15-DK × 3（1 发 + 2 收）",
        "USB 数据线 × 3",
        "逻辑分析仪（≥ 20MS/s）",
        "杜邦线 / 测试探针",
    ]),
    ("软件 Software", BLUE2, [
        "nRF Connect SDK v3.4.0",
        "Zephyr RTOS 4.4",
        "本工程固件 zephyr.hex",
        "编译：west build",
    ]),
    ("工具 Tools", CYAN, [
        "nrfutil（烧录 / 复位）",
        "串口终端（PuTTY / minicom）",
        "逻辑分析仪软件（Saleae 等）",
        "Python（文档/图表生成）",
    ]),
]
x = 0.5
for t, c, items in cols:
    add_rect(s, x, 1.7, 3.0, 4.7, WHITE, line=c)
    add_rect(s, x, 1.7, 3.0, 0.7, c, shape=MSO_SHAPE.RECTANGLE)
    add_shape_text(add_rect(s, x, 1.7, 3.0, 0.7, c, shape=MSO_SHAPE.RECTANGLE), t, 14, WHITE)
    add_text(s, x + 0.25, 2.6, 2.6, 3.6, "\n".join("• " + it for it in items), 12, INK)
    x += 3.15
# 实测搭建插图面板
add_rect(s, x, 1.7, 3.0, 4.7, WHITE, line=BLUE2)
s.shapes.add_picture(f"{IMG}/Development_board_and_logic_an_2026-09-15T04-52-53.png", I(x + 0.15), I(1.85), width=I(2.7))
add_text(s, x + 0.1, 4.7, 2.8, 1.4, "实测搭建：大脑主控板\n+ 关节节点 + 逻辑分析仪", 12, GRAY, align=PP_ALIGN.CENTER)

# ================= 8 波形示意 =================
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
add_text(s, 0.55, 0.35, 12, 0.8, "演示效果：三通道脉冲边沿对齐", 28, WHITE, True)
add_rect(s, 0.55, 1.15, 2.2, 0.1, CYAN)
chan = [("TX  发送端", CYAN), ("RX1 接收端", BLUE2), ("RX2 接收端", BLUE2)]
base_y = [2.3, 3.6, 4.9]
for i, (name, c) in enumerate(chan):
    add_text(s, 0.7, base_y[i] - 0.42, 2.0, 0.4, name, 16, WHITE, True)
    add_rect(s, 2.6, base_y[i], 9.6, 0.05, GRAY, shape=MSO_SHAPE.RECTANGLE)
    for px in [4.2, 6.6, 9.0]:
        add_rect(s, px, base_y[i] - 0.35, 0.22, 0.7, c, shape=MSO_SHAPE.RECTANGLE)
for px in [4.31, 6.71, 9.11]:
    add_rect(s, px, 1.5, 0.02, 4.1, (0x40, 0x60, 0x80), shape=MSO_SHAPE.RECTANGLE)
add_text(s, 0.7, 6.0, 11.5, 0.8, "三通道脉冲上升沿基本重合 —— 边沿时间差即同步误差（µs 级）", 15, CYAN, align=PP_ALIGN.CENTER)

# ================= 9 实测数据 =================
s = prs.slides.add_slide(BLANK)
bg(s, LIGHT)
title_bar(s, "实测数据", "MEASURED RESULT")
add_text(s, 0.9, 1.9, 5.0, 1.6, "0 µs", 72, BLUE, True)
add_text(s, 0.9, 3.5, 5.0, 0.8, "时间戳一致误差", 18, INK, True)
sp = add_rect(s, 6.4, 1.9, 6.2, 3.4, NAVY)
add_text(s, 6.7, 2.1, 5.6, 2.6,
         "TX: Sent SDU counter 205600\n     timestamp 1028163407 us\n"
         "RX: Recv SDU counter 205600\n     shared_ts 1028163407 us\n"
         "→ 完全一致", 14, CYAN)
add_text(s, 6.4, 5.5, 6.2, 1.0, "脉冲边沿三通道对齐：µs 级", 14, GRAY, align=PP_ALIGN.CENTER)

# ================= 10 技术优势 =================
s = prs.slides.add_slide(BLANK)
bg(s, LIGHT)
title_bar(s, "技术优势", "ADVANTAGES")
adv = [
    ("微秒级精度", [
        "controller-timed 硬件呈现（GRTC/DPPI/GPIOTE）",
        "时钟伺服 + 漂移外推，消除陈旧误差",
        "实测时间戳一致误差 0µs",
    ]),
    ("长期稳定可靠", [
        "在线估计晶振漂移（ppm），持续跟踪",
        "失同步 holdover 按最后速率滑行",
        "32 位时间戳回绕正确处理",
    ]),
    ("轻量易扩展", [
        "一对多广播，无连接、无回传",
        "接收端数量可扩展",
        "O(1) 伺服、无动态内存、200Hz 开销可忽略",
    ]),
    ("硬件级确定性", [
        "GPIO 呈现免软件线程调度抖动",
        "SET/CLR 绝对电平、幂等自愈",
        "中位数估计，对丢包/重传鲁棒",
    ]),
]
pos = [(0.7, 1.7), (6.85, 1.7), (0.7, 4.3), (6.85, 4.3)]
for (t, bullets), (px, py) in zip(adv, pos):
    add_rect(s, px, py, 5.8, 2.45, WHITE, line=BLUE)
    add_rect(s, px, py, 5.8, 0.55, BLUE, shape=MSO_SHAPE.RECTANGLE)
    add_shape_text(add_rect(s, px, py, 5.8, 0.55, BLUE, shape=MSO_SHAPE.RECTANGLE), t, 15, WHITE)
    add_text(s, px + 0.3, py + 0.72, 5.2, 1.6, "\n".join("• " + b for b in bullets), 12.5, INK)

# ================= 11 应用场景 =================
scenes = [
    ("机器人协同", "机器人大脑与四肢\n关节时钟同步", "Whole_body_coordinated_motion__2026-09-15T04-52-57.png"),
    ("工业控制", "多轴运动控制\n分布式 PLC 同步", "Industrial_multi_axis_motion_c_2026-09-15T05-00-02.png"),
    ("无线传感网络", "分布式 IoT 节点\n同步采集", "Wireless_sensor_network_of_dis_2026-09-15T05-00-01.png"),
    ("声学阵列", "麦克风阵列\n波束成形", "Microphone_array_beamforming_w_2026-09-15T05-00-03.png"),
    ("振动监测", "旋转机械多通道\n振动同步采集", "Vibration_monitoring_of_rotati_2026-09-15T05-00-07.png"),
    ("电力系统", "智能电网\nPMU 相量测量", "Smart_grid_electrical_substati_2026-09-15T05-00-04.png"),
    ("通信基站", "5G TDD\n时间同步", "Telecom_base_station_tower_wit_2026-09-15T05-00-07.png"),
    ("金融交易", "低延迟行情\n时间戳同步", "Financial_trading_system_with__2026-09-15T05-00-11.png"),
    ("自动驾驶", "多传感器融合\nLiDAR / 相机 / IMU", "Autonomous_vehicle_sensor_fusi_2026-09-15T05-00-06.png"),
    ("音视频同步", "多相机多麦克风\n音画同步", "Multi_camera_and_microphone_st_2026-09-15T05-00-12.png"),
    ("科学测量", "射电望远镜阵列\n分布式测量", "Radio_telescope_array_with_syn_2026-09-15T05-00-07.png"),
    ("医疗设备", "多设备\n同步监测", "Medical_monitoring_devices_syn_2026-09-15T05-00-09.png"),
]
PER = 6
for pg in range(2):
    s = prs.slides.add_slide(BLANK)
    bg(s, LIGHT)
    title_bar(s, "应用场景" if pg == 0 else "应用场景（续）",
              ("SCENARIOS · 所有需要时钟 / 信号同步的领域" if pg == 0
               else "SCENARIOS · 更多同步应用领域"))
    chunk = scenes[pg * PER:(pg + 1) * PER]
    x0, y0, cw, ch, gx, gy = 0.55, 1.7, 4.0, 2.62, 0.15, 0.22
    for i, (t, d, fn) in enumerate(chunk):
        r, c = divmod(i, 3)
        x = x0 + c * (cw + gx)
        y = y0 + r * (ch + gy)
        add_rect(s, x, y, cw, ch, CARD, line=BLUE2)
        add_rect(s, x, y, 0.14, ch, BLUE, shape=MSO_SHAPE.RECTANGLE)
        s.shapes.add_picture(f"{IMG}/{fn}", I(x + 0.22), I(y + 0.48), width=I(1.65))
        add_text(s, x + 2.0, y + 0.32, 1.9, 0.45, t, 14, NAVY, True)
        add_text(s, x + 2.0, y + 0.8, 1.9, 1.6, d, 10.5, GRAY)

# ================= 12 总结 =================
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
deco_dots(s, 6)
add_text(s, 0.7, 1.7, 6.0, 1.0, "总结", 40, WHITE, True)
add_text(s, 0.7, 2.9, 6.0, 2.8,
         "基于 BLE BIS + 控制器时间戳 + 时钟伺服\n\n实现微秒级多节点时间同步\n\n"
         "实测时间戳一致（0µs）· 边沿 µs 级对齐", 18, CYAN)
add_rect(s, 0.75, 6.0, 2.6, 0.06, CYAN, shape=MSO_SHAPE.RECTANGLE)
add_text(s, 0.7, 6.25, 6.0, 0.6, "感谢观看 · 欢迎交流", 16, WHITE)
add_rect(s, 6.9, 1.1, 6.1, 4.15, None, line=CYAN)
s.shapes.add_picture(f"{IMG}/Brain_to_joint_clock_synchroni_2026-09-15T04-52-51.png", I(7.0), I(1.2), width=I(5.9))

prs.save(OUT)
print("PPT 已重新生成：", OUT)
