# -*- coding: utf-8 -*-
"""一页版产品介绍：BIS 时间同步方案（参考 temp.pptx 绿色产品页风格）。"""

from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

OUT = "/home/devin/ncs/lim_work/iso_time_sync/BIS时间同步方案_一页版.pptx"
IMG = "/home/devin/ncs/lim_work/iso_time_sync/scene_images"

# ---- temp.pptx 绿色主题配色 ----
GREEN  = (0x41, 0xC3, 0x63)   # accent1
GREEN2 = (0x8D, 0xDB, 0xA1)   # dk2
AQUA   = (0x80, 0xDC, 0xEB)   # accent2
YELLOW = (0xFF, 0xD1, 0x00)   # accent3
LGREEN = (0xBF, 0xEB, 0xC9)   # lt2
DGRAY  = (0x4D, 0x4D, 0x4D)   # accent6
MGRAY  = (0x91, 0x91, 0x91)   # accent5
WHITE  = (0xFF, 0xFF, 0xFF)
BLACK  = (0x00, 0x00, 0x00)

prs = Presentation()
prs.slide_width = I(13.333)
prs.slide_height = I(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = C(*WHITE)

def set_font(run, size=14, color=BLACK, bold=False, name="微软雅黑"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = C(*color)
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', name)

def add_rect(x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = s.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = C(*fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = C(*line); sp.line.width = Pt(1.2)
    sp.shadow.inherit = False
    return sp

def add_text(x, y, w, h, text, size=14, color=BLACK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(5)
        r = p.add_run(); r.text = ln
        set_font(r, size=size, color=color, bold=bold)
    return tb

def shape_text(sp, text, size=14, color=WHITE, bold=True):
    tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(8); tf.margin_right = Pt(6)
    tf.margin_top = Pt(3); tf.margin_bottom = Pt(3)
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = ln
        set_font(r, size=size, color=color, bold=bold)

def section_bar(x, y, w, title):
    add_rect(x, y, w, 0.5, GREEN, shape=MSO_SHAPE.RECTANGLE)
    shape_text(add_rect(x, y, w, 0.5, GREEN, shape=MSO_SHAPE.RECTANGLE), title, 15, WHITE)

# ---- 顶部标题 ----
add_text(0.5, 0.08, 12.3, 0.65, "nRF54L15 BLE BIS 时间同步方案", 28, BLACK, True)
add_text(0.5, 0.66, 12.3, 0.35, "Bluetooth Low Energy · Broadcast Isochronous Stream · Time Synchronization", 12, MGRAY)
add_rect(0.5, 0.95, 2.2, 0.05, GREEN, shape=MSO_SHAPE.RECTANGLE)

# ---- 左侧：主视觉图 + Overview ----
add_rect(0.4, 1.1, 6.35, 4.23, None, line=LGREEN)
s.shapes.add_picture(f"{IMG}/Clock_synchronization_of_multi_2026-09-15T05-24-38.png", I(0.4), I(1.1), width=I(6.35))
section_bar(0.4, 5.45, 6.35, "Overview")
add_text(0.6, 6.08, 6.0, 1.3,
    "• BLE 5.2 BIS 一对多广播：1 发送 + 2 接收共享同一时间基准\n"
    "• 发送端在 SDU 嵌入控制器时间戳，接收端时钟伺服换算本地 GRTC\n"
    "• 长期 µs 级时间戳同步，三方对同一事件打出相同时间戳", 13, BLACK)

# ---- 右侧：Key Features ----
section_bar(7.0, 1.1, 5.95, "Key Features")
add_text(7.25, 1.75, 5.5, 2.9,
    "• 微秒级精度 —— 实测时间戳一致误差 0 µs\n"
    "• 硬件级确定性 —— GRTC/DPPI/GPIOTE 驱动 P1.10 1µs 脉冲\n"
    "• 长期稳定 —— 在线估计晶振漂移(ppm)，失同步自动恢复\n"
    "• 轻量易扩展 —— 无连接一对多广播，接收端数量可扩展\n"
    "• 高效 —— O(1) 时钟伺服，200Hz 更新，开销可忽略\n"
    "• 鲁棒 —— 中位数估计，对丢包/重传免疫", 12.5, BLACK)

# ---- 右侧：主要应用场景 ----
section_bar(7.0, 4.8, 5.95, "主要应用场景  Target Application")
add_text(7.25, 5.45, 2.8, 1.9,
    "• 机器人协同（大脑-关节）\n"
    "• 工业控制（多轴运动）\n"
    "• 无线传感网络（分布式）\n"
    "• 声学阵列（波束成形）", 12.5, BLACK)
add_text(10.15, 5.45, 2.7, 1.9,
    "• 金融交易（低延迟戳）\n"
    "• 自动驾驶（多传感器）\n"
    "• 电力系统（PMU 相量）\n"
    "• 科学测量（射电阵列）", 12.5, BLACK)

prs.save(OUT)
print("一页版 PPT 已生成：", OUT)
